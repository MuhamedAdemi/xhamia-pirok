"""
Importi i shtëpive nga Excel (fleta "Shtëpitë", formati i backup.KOLONAT_SHTEPITE).

Për çdo rresht:
  - "Borxh nga viti" X   → antarësia llogaritet nga X; vitet para X konsiderohen të paguara
  - "Paguar deri viti" Y → (pa borxh) e paguar deri në Y
Shtëpitë gjenden sipas emër + prind + mbiemër. Kur i njëjti emër ekziston disa herë (p.sh. dy
shtëpi "Gajur Shemsedin Aliu"), rreshtat lidhen me shtëpitë ekzistuese sipas radhës; të tjerat krijohen.
"""
import re
from collections import defaultdict

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from .importi_pdf import celesi
from .models import Shtepia


def _pastro_koken(k):
    return re.sub(r'\s+', ' ', str(k or '').strip().lower().replace('ë', 'e'))


_KOLONAT = {
    'nr.': 'nr', 'nr': 'nr', 'emri': 'emri', 'emri i prindit': 'prindi', 'prindi': 'prindi',
    'mbiemri': 'mbiemri', 'telefoni': 'telefoni', 'telefon': 'telefoni',
    'paguar deri viti': 'paguar_deri', 'borxh nga viti': 'borxh_nga',
}


def _viti(v):
    if v in (None, ''):
        return None
    try:
        viti = int(float(str(v).strip()))
    except ValueError:
        raise ValueError(f'"{v}" nuk është vit.')
    if not 1950 <= viti <= 2100:
        raise ValueError(f'Viti {viti} nuk duket i saktë.')
    return viti


def lexo_excel(skedari):
    wb = load_workbook(skedari, read_only=True, data_only=True)
    ws = wb['Shtëpitë'] if 'Shtëpitë' in wb.sheetnames else wb.worksheets[0]
    rreshtat = ws.iter_rows(values_only=True)
    koka = [_KOLONAT.get(_pastro_koken(k)) for k in next(rreshtat, [])]
    if 'emri' not in koka or 'mbiemri' not in koka:
        raise ValueError('Excel-i duhet të ketë së paku kolonat "Emri" dhe "Mbiemri" në rreshtin e parë.')

    rezultati, gabime = [], []
    for i, vlerat in enumerate(rreshtat, start=2):
        r = {k: v for k, v in zip(koka, vlerat) if k}
        emri = str(r.get('emri') or '').strip()
        mbiemri = str(r.get('mbiemri') or '').strip()
        if not emri and not mbiemri:
            continue
        try:
            rezultati.append({
                'rreshti': i, 'nr': str(r.get('nr') or '').strip(),
                'emri': emri, 'prindi': str(r.get('prindi') or '').strip(), 'mbiemri': mbiemri,
                'telefoni': str(r.get('telefoni') or '').strip(),
                'paguar_deri': _viti(r.get('paguar_deri')), 'borxh_nga': _viti(r.get('borxh_nga')),
            })
        except ValueError as e:
            gabime.append(f'Rreshti {i} ({emri} {mbiemri}): {e}')
    return rezultati, gabime


def importo_excel(skedari, perdoruesi, vetem_prove=False):
    rreshtat, gabime = lexo_excel(skedari)
    viti_aktual = timezone.now().year
    raporti = {
        'gjithsej_rreshta': len(rreshtat), 'te_reja': 0, 'perditesuara': 0,
        'paguar': 0, 'borxh_vetem_aktual': 0, 'borxh_shume_vite': 0, 'pa_gjendje': 0,
        'vite_borxh': 0, 'gabime': gabime, 'nr_nga': None, 'nr_deri': None,
    }
    if gabime:
        return raporti

    ekzistuese = defaultdict(list)
    for sh in Shtepia.objects.order_by('nr_shtepise', 'pk'):
        ekzistuese[celesi(sh.emri_kryefamiljarit, sh.emri_prindit, sh.mbiemri_kryefamiljarit)].append(sh)
    numrat_zene = set(Shtepia.objects.values_list('nr_shtepise', flat=True))
    nr_radhes = 1   # mbush numrat e lirë (p.sh. të shtëpive të fshira) para se të shkojë më tej

    with transaction.atomic():
        te_reja = []
        for r in rreshtat:
            # Gjendja e antarësisë sipas rreshtit
            if r['borxh_nga']:
                fillimi, paguar_deri = r['borxh_nga'], r['borxh_nga'] - 1
                vite = max(0, viti_aktual - r['borxh_nga'] + 1)
                raporti['vite_borxh'] += vite
                raporti['borxh_shume_vite' if vite > 1 else 'borxh_vetem_aktual'] += 1
            elif r['paguar_deri']:
                fillimi, paguar_deri = min(r['paguar_deri'], viti_aktual), r['paguar_deri']
                raporti['paguar'] += 1
            else:
                fillimi = paguar_deri = None
                raporti['pa_gjendje'] += 1

            lista = ekzistuese.get(celesi(r['emri'], r['prindi'], r['mbiemri']))
            if lista:
                sh = lista.pop(0)
                ndryshime = {}
                if fillimi is not None:
                    ndryshime.update(viti_fillimit_antaresise=fillimi, paguar_deri_viti=paguar_deri)
                if r['telefoni'] and not sh.nr_telefoni_kryesor:
                    ndryshime['nr_telefoni_kryesor'] = r['telefoni']
                if ndryshime:
                    Shtepia.objects.filter(pk=sh.pk).update(**ndryshime)
                raporti['perditesuara'] += 1
                continue

            nr = r['nr'].zfill(4) if r['nr'].isdigit() else ''
            if not nr or nr in numrat_zene:
                while f'{nr_radhes:04d}' in numrat_zene:
                    nr_radhes += 1
                nr = f'{nr_radhes:04d}'
            numrat_zene.add(nr)
            te_reja.append(Shtepia(
                nr_shtepise=nr, emri_kryefamiljarit=r['emri'], emri_prindit=r['prindi'],
                mbiemri_kryefamiljarit=r['mbiemri'], nr_telefoni_kryesor=r['telefoni'],
                viti_fillimit_antaresise=fillimi or viti_aktual, paguar_deri_viti=paguar_deri,
                regjistruar_nga=perdoruesi,
            ))

        Shtepia.objects.bulk_create(te_reja, batch_size=200)
        raporti['te_reja'] = len(te_reja)
        if te_reja:
            numrat = sorted(sh.nr_shtepise for sh in te_reja)
            raporti['nr_nga'], raporti['nr_deri'] = numrat[0], numrat[-1]
        if vetem_prove:
            transaction.set_rollback(True)
    return raporti
